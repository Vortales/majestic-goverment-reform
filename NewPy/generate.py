# -*- coding: utf-8 -*-
"""Generate ACT. <num>.docx bills from content_A..E specs,
strictly following the reference layout of ACT. 000.docx:

  ACT. NNN                                  (centred, Century Schoolbook 37 pt bold)
  IN THE SENATE OF THE STATE OF SAN ANDREAS (centred, Century Schoolbook 15 pt)
  MARCH 00, 2026                            (centred, Century Schoolbook 15 pt)
  ЗАКОНОПРОЕКТ                              (centred, Century Schoolbook 30 pt)
  о внесении изменений в действующий закон  (centred, Times New Roman 14 pt)
  SECTION 1. СУТЬ ПРАВКИ.                   (Times New Roman 12 pt bold)
  <essence>                                 (Times New Roman 12 pt bold)
  SECTION. 2. ВНЕСЕНИЕ ПРАВКИ В <law>       (Times New Roman 12 pt bold)
  + table 2 columns:
      | настоящая редакция | предложенный вариант |   (header row, centred)
      | full current article, changed fragment shaded F4CCCC (light red)
      | full proposed article, changed fragment shaded D9EAD3 (light green)
      ... one row per article

Markers produced by content_*.py:
  {{...}} — old fragment (red) in the "настоящая редакция" column
  [[...]] — new fragment (green) in the "предложенный вариант" column
"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "ACT. 000.docx")

RED_FILL = "F4CCCC"
GREEN_FILL = "D9EAD3"
BODY_FONT = "Times New Roman"
HEAD_FONT = "Century Schoolbook"
TEXT_COLOR = "333333"

TOKEN = re.compile(r"(\{\{.*?\}\}|\[\[.*?\]\])", re.S)

# reverse schema order of CT_RPr children that may legally precede w:shd
_SHD_PRECEDERS = [
    "w:bdr", "w:effect", "w:u", "w:highlight", "w:szCs", "w:sz",
    "w:position", "w:kern", "w:spacing", "w:color", "w:bCs", "w:b",
    "w:rFonts", "w:rStyle",
]


def set_font(run, name, size_pt, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def set_shd(run, fill):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    for tag in _SHD_PRECEDERS:
        el = rPr.find(qn(tag))
        if el is not None:
            el.addnext(shd)
            return
    rPr.append(shd)


def copy_ppr(paragraph, template_ppr):
    if template_ppr is None:
        return
    old = paragraph._p.find(qn("w:pPr"))
    if old is not None:
        paragraph._p.remove(old)
    paragraph._p.insert(0, copy.deepcopy(template_ppr))


def split_marked(text):
    """-> [[ [line_text, marked], ... ], ...] paragraph-wise."""
    paras = [[]]
    for part in TOKEN.split(text):
        if not part:
            continue
        if (part.startswith("{{") and part.endswith("}}")) or \
           (part.startswith("[[") and part.endswith("]]")):
            inner, marked = part[2:-2], True
        else:
            inner, marked = part, False
        lines = inner.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                paras.append([])
            if line:
                paras[-1].append((line, marked))
    return paras


def load_template_parts():
    doc = Document(TEMPLATE)
    body = doc.element.body
    # capture paragraph pPr in document order (body-level paragraphs)
    p_elems = [el for el in body if el.tag == qn("w:p")]
    pPr_list = [el.find(qn("w:pPr")) for el in p_elems]
    tbl = doc.tables[0]
    tblPr = copy.deepcopy(tbl._tbl.tblPr)
    tblGrid = copy.deepcopy(tbl._tbl.find(qn("w:tblGrid")))
    header_tr = copy.deepcopy(tbl.rows[0]._tr)
    data_tr = copy.deepcopy(tbl.rows[1]._tr)
    return doc, pPr_list, tblPr, tblGrid, header_tr, data_tr


def clear_body(doc):
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)
    return sectPr


def add_paragraph(doc, pPr_template, runs, align=None):
    p = doc.add_paragraph()
    copy_ppr(p, pPr_template)
    if align is not None:
        p.alignment = align
    for text, fmt in runs:
        r = p.add_run(text)
        set_font(r, fmt.get("font", BODY_FONT), fmt.get("size", 12),
                 fmt.get("bold", False), fmt.get("color"))
    return p


def build_header(doc, pPr_list, spec):
    # indices follow the paragraph order of ACT. 000.docx
    add_paragraph(doc, pPr_list[0],
                  [("ACT. %d" % spec["num"],
                    {"font": HEAD_FONT, "size": 37, "bold": True, "color": TEXT_COLOR})],
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[1], [])
    add_paragraph(doc, pPr_list[2],
                  [("IN THE SENATE OF THE STATE OF SAN ANDREAS",
                    {"font": HEAD_FONT, "size": 15, "color": TEXT_COLOR})],
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[3],
                  [("MARCH 00, 2026",
                    {"font": HEAD_FONT, "size": 15, "color": TEXT_COLOR})],
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[4], [])
    add_paragraph(doc, pPr_list[5],
                  [("ЗАКОНОПРОЕКТ",
                    {"font": HEAD_FONT, "size": 30, "color": TEXT_COLOR})],
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[6],
                  [("о внесении изменений в действующий закон",
                    {"font": BODY_FONT, "size": 14, "color": TEXT_COLOR})],
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[7], [], align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, pPr_list[8],
                  [("SECTION 1.", {"size": 12, "bold": True, "color": TEXT_COLOR}),
                   (" ", {"size": 12, "color": TEXT_COLOR}),
                   ("СУТЬ ПРАВКИ.", {"size": 12, "bold": True, "color": TEXT_COLOR})])
    add_paragraph(doc, pPr_list[9],
                  [(spec["essence"], {"size": 12, "bold": True, "color": TEXT_COLOR})])
    add_paragraph(doc, pPr_list[10],
                  [("SECTION. 2. ВНЕСЕНИЕ ПРАВКИ В " + spec["law"],
                    {"size": 12, "bold": True, "color": TEXT_COLOR})])
    add_paragraph(doc, pPr_list[11], [])


def build_table(doc, spec, tblPr, tblGrid, header_tr, data_tr):
    table = doc.add_table(rows=0, cols=2)
    tbl = table._tbl
    # swap in the template's properties / grid
    for old in list(tbl):
        if old.tag in (qn("w:tblPr"), qn("w:tblGrid")):
            tbl.remove(old)
    tbl.insert(0, copy.deepcopy(tblPr))
    tbl.insert(1, copy.deepcopy(tblGrid))
    # header row straight from the template ("настоящая редакция" | "предложенный вариант")
    tbl.append(copy.deepcopy(header_tr))

    for art in spec["articles"]:
        tbl.append(copy.deepcopy(data_tr))
        row = table.rows[-1]
        for col, key in ((0, "current"), (1, "proposed")):
            cell = row.cells[col]
            fill_cell(cell, art[key], RED_FILL if col == 0 else GREEN_FILL)
    return table


def fill_cell(cell, text, fill):
    first = cell.paragraphs[0]
    pPr_tmpl = copy.deepcopy(first._p.find(qn("w:pPr")))
    # strip everything but pPr of the first paragraph, drop the rest
    for child in list(first._p):
        if child.tag != qn("w:pPr"):
            first._p.remove(child)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

    paras = split_marked(text)
    for pi, runs in enumerate(paras):
        p = first if pi == 0 else cell.add_paragraph()
        if pi > 0:
            copy_ppr(p, pPr_tmpl)
        for txt, marked in runs:
            r = p.add_run(txt)
            set_font(r, BODY_FONT, 12)
            if marked:
                set_shd(r, fill)


def verify(path, spec):
    """Compare the generated document against the ACT. 000 reference structure."""
    problems = []
    d = Document(path)
    paras = [p for p in d.paragraphs]
    if not paras or paras[0].text != "ACT. %d" % spec["num"]:
        problems.append("title line mismatch")
    else:
        r = paras[0].runs[0]
        if r.font.name != HEAD_FONT or r.font.size.pt != 37 or not r.font.bold:
            problems.append("title font mismatch")
        if d.paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.CENTER:
            problems.append("title not centred")
    if len(d.tables) != 1:
        problems.append("expected exactly 1 table")
        return problems
    t = d.tables[0]
    if len(t.rows) != len(spec["articles"]) + 1:
        problems.append("row count != articles + 1")
    hdr = [t.rows[0].cells[0].text.strip(), t.rows[0].cells[1].text.strip()]
    if hdr != ["настоящая редакция", "предложенный вариант"]:
        problems.append("table header labels mismatch: %r" % (hdr,))
    n_red = n_green = 0
    body = d.element.body
    for shd in body.iter(qn("w:shd")):
        f = shd.get(qn("w:fill"))
        if f == RED_FILL:
            n_red += 1
        elif f == GREEN_FILL:
            n_green += 1
    if n_red == 0 and n_green == 0:
        problems.append("no highlighted fragments at all")
    # no raw markers must survive into the document
    full = "\n".join(p.text for p in d.paragraphs)
    for t_ in d.tables:
        for row in t_.rows:
            for c in row.cells:
                full += "\n" + c.text
    for lit in ("{{", "}}", "[[", "]]", "…"):
        if lit in full:
            problems.append("literal %r left in document" % lit)
    # fonts inside the table
    cell_run = None
    for row in t.rows[1:]:
        for c in row.cells:
            for p in c.paragraphs:
                if p.runs:
                    cell_run = p.runs[0]
                    break
            if cell_run:
                break
        if cell_run:
            break
    if cell_run is not None:
        if cell_run.font.name != BODY_FONT or cell_run.font.size.pt != 12:
            problems.append("cell font mismatch (%s/%s)" %
                            (cell_run.font.name, cell_run.font.size))
    return problems


def main():
    sys.path.insert(0, HERE)
    import content_A, content_B, content_C, content_D, content_E
    specs = (content_A.SPECS_A + content_B.SPECS_B + content_C.SPECS_C +
             content_D.SPECS_D + content_E.SPECS_E)

    ok = True
    for spec in specs:
        doc, pPr_list, tblPr, tblGrid, header_tr, data_tr = load_template_parts()
        clear_body(doc)
        build_header(doc, pPr_list, spec)
        build_table(doc, spec, tblPr, tblGrid, header_tr, data_tr)
        # keep section properties as the last body element
        body = doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        if sectPr is not None:
            body.remove(sectPr)
            body.append(sectPr)
        out = os.path.join(HERE, "ACT. %d.docx" % spec["num"])
        doc.save(out)
        problems = verify(out, spec)
        status = "OK " if not problems else "FAIL"
        if problems:
            ok = False
        print("%s ACT. %d.docx  (%d article rows%s)%s" %
              (status, spec["num"], len(spec["articles"]),
               ", %d red / %d green marks" % (
                   sum(a["current"].count("{{") for a in spec["articles"]),
                   sum(a["proposed"].count("[[") for a in spec["articles"]))
               if not problems else "",
               ("  -> " + "; ".join(problems)) if problems else ""))
    print("\n%s" % ("ALL DOCUMENTS GENERATED AND VERIFIED" if ok else "SOME DOCUMENTS HAVE PROBLEMS"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
