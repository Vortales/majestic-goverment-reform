#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Проверка собранного .docx: структура, цвета, полнота текста."""
import sys, os
from docx import Document
from docx.shared import RGBColor

RED = RGBColor(0xC0,0,0); GREEN = RGBColor(0,0x7A,0)

def check(path, verbose=False):
    d = Document(path)
    print(f"### {os.path.basename(path)}")
    # заголовочная часть
    head = [p.text for p in d.paragraphs[:8]]
    for h in head: print("   |", h[:100])
    tabs = d.tables
    print(f"   таблиц: {len(tabs)}")
    for ti, t in enumerate(tabs):
        print(f"   таблица {ti}: строк={len(t.rows)} колонок={len(t.columns)}")
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                nred = ngreen = nplain = 0
                for p in cell.paragraphs:
                    for r in p.runs:
                        c = r.font.color and r.font.color.rgb
                        if c == RED: nred += len(r.text)
                        elif c == GREEN: ngreen += len(r.text)
                        else: nplain += len(r.text)
                txt = cell.text
                if verbose:
                    print(f"     [{ri},{ci}] символов={len(txt)} red={nred} green={ngreen} plain={nplain}")
                    print("       " + txt[:400].replace("\n"," / "))
                else:
                    ell = txt.count("…") + txt.count("...")
                    print(f"     [{ri},{ci}] len={len(txt):5d} red={nred:4d} green={ngreen:4d} многоточий={ell}")
                    if ell: print("        !!! МНОГОТОЧИЕ:", [s for s in txt.split("\n") if "…" in s or "..." in s][:2])
    print()

if __name__ == "__main__":
    v = "-v" in sys.argv
    for a in sys.argv[1:]:
        if a == "-v": continue
        p = a if a.endswith(".docx") else f"ACT/ACT.{a}.docx"
        check(p, v)
