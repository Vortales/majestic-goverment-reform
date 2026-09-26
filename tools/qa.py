#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""QA собранного .docx: артефакты разметки, многоточия, пустые ячейки, ссылки на ACT/проекты."""
import sys, os, re, glob
from docx import Document

BAD = [
    (re.compile(r'"\s*\n\s*"'), "склейка строк (кавычка в конце строки)"),
    (re.compile(r'"\s*\+\s*"'), "склейка строк (+ между кавычками)"),
    (re.compile(r"\{\+|\+\}|\{-|-\}"), "нераскрытый маркер {- -} / {+ +}"),
    (re.compile(r"…|\.\.\."), "многоточие (усечение текста)"),
    (re.compile(r"ACT\s*\.\s*\d", re.I), "ссылка на номер акта (ACT.)"),
    (re.compile(r"[Пп]роект\w*\s*№"), "ссылка на номер проекта"),
    (re.compile(r"\bсм\.\s*(выше|ниже|проект)"), "служебная отсылка"),
]

def qa(path):
    d = Document(path)
    issues = []
    def scan(txt, where):
        for rx, name in BAD:
            for m in rx.finditer(txt):
                ctx = txt[max(0, m.start()-45):m.end()+45].replace("\n", " / ")
                issues.append(f"{where}: {name}: …{ctx}…")
    for i, p in enumerate(d.paragraphs):
        if i == 0:
            continue  # заголовок документа «ACT. NNN» — обязателен по образцу
        scan(p.text, f"параграф {i}")
    for ti, t in enumerate(d.tables):
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                txt = cell.text
                if ri > 0 and not txt.strip():
                    issues.append(f"таблица {ti} строка {ri} колонка {ci}: ПУСТАЯ ЯЧЕЙКА")
                scan(txt, f"таблица {ti} строка {ri} колонка {ci}")
    return issues

if __name__ == "__main__":
    files = sys.argv[1:] or sorted(glob.glob("ACT/ACT.*.docx"))
    total = 0
    for f in files:
        if not f.endswith(".docx"):
            f = f"ACT/ACT.{f}.docx"
        iss = qa(f)
        total += len(iss)
        print(f"### {os.path.basename(f)}: {'OK' if not iss else str(len(iss)) + ' замечаний'}")
        for i in iss[:25]:
            print("   -", i)
    print(f"\nИТОГО замечаний: {total}")
