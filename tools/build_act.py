#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор документов ACT.<номер>.docx в формате «Приложения №2 — о внесении изменений
в действующий закон».

Формат строки таблицы:
  R(src=(акт, "номер статьи"), red=["фрагмент 1", ...], new="текст с {+зелёным+}")
  R(old="полный текст статьи с {-красным-}", new="полный текст статьи с {+зелёным+}")
"""
import os, re, sys, importlib

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "Предлагаемое на сенат законодательство")
OUT = os.path.join(BASE, "ACT")

ZW = "\u200b"
STRUCT = re.compile(r"^\s*(Особенная часть|Общая часть|Раздел|РАЗДЕЛ|Глава|ГЛАВА|Уровень|Глава\s+\d)\b")

# ---------------------------------------------------------------- источник ----
_cache = {}

def load(act):
    if act not in _cache:
        p = os.path.join(SRC, f"{act}.txt")
        t = open(p, encoding="utf-8-sig").read()
        t = t.replace("\r\n", "\n").replace("\r", "\n").replace(ZW, "")
        _cache[act] = t
    return _cache[act]

def split_articles(act):
    t = load(act)
    if act == 163:
        pat = re.compile(r"^\s*(\d+(?:\.\d+)+)\s", re.M)
    else:
        pat = re.compile(r"^\s*Статья\.?\s+([0-9]+(?:\.[0-9]+)*)", re.M)
    ms = list(pat.finditer(t))
    out = []
    for i, m in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(t)
        out.append((m.group(1), t[m.start():end]))
    return out

def clean(body):
    """Убирает «хвост» статьи: заголовки разделов/глав, идущие после текста статьи."""
    lines = body.split("\n")
    keep = [lines[0]]
    for ln in lines[1:]:
        if STRUCT.match(ln):
            break
        keep.append(ln)
    txt = "\n".join(keep)
    txt = re.sub(r"[ \t]+\n", "\n", txt)
    txt = re.sub(r"\n{3,}", "\n\n", txt)
    return txt.strip()

def get_article(act, num):
    for n, b in split_articles(act):
        if n == num:
            return clean(b)
    raise KeyError(f"ACT {act}: статья {num} не найдена")

def norm(s):
    """Нормализация для поиска фрагментов (пробелы, кавычки, тире)."""
    s = s.replace(ZW, "").replace("\u00a0", " ")
    s = s.replace("\r", " ").replace("\n", " ")
    s = re.sub(r"\s+", " ", s)
    s = s.replace("«", '"').replace("»", '"').replace("„", '"').replace("“", '"').replace("”", '"')
    s = s.replace("—", "-").replace("–", "-").replace("ё", "е").replace("Ё", "Е")
    return s.strip()

def mark_red(text, fragments):
    """Разбивает текст на сегменты [(фрагмент, 'red'|''), ...]."""
    if not fragments:
        return [(text, "")]
    spans = []
    for fr in fragments:
        i = find_span(text, fr)
        if i is None:
            raise ValueError(f"ФРАГМЕНТ НЕ НАЙДЕН В СТАТЬЕ:\n  ищем: {fr[:120]!r}")
        spans.append((i, i + len(fr)))
    spans.sort()
    merged = []
    for s, e in spans:
        if merged and s <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], e))
        else:
            merged.append((s, e))
    segs, pos = [], 0
    for s, e in merged:
        if s > pos:
            segs.append((text[pos:s], ""))
        segs.append((text[s:e], "red"))
        pos = e
    if pos < len(text):
        segs.append((text[pos:], ""))
    return segs

def find_span(text, fr):
    """Возвращает позицию фрагмента `fr` в `text` с учётом нормализации пробелов/кавычек."""
    nf = norm(fr)
    if not nf:
        return None
    buf = []
    for ch in text:
        c = ch
        if c == "\u00a0":
            c = " "
        if c in "\r\n\t":
            c = " "
        if c in "\u00ab\u00bb\u201e\u201c\u201d":
            c = '"'
        if c in "\u2014\u2013":
            c = "-"
        if c == "\u0451":
            c = "\u0435"
        if c == "\u0401":
            c = "\u0415"
        buf.append(c)
    collapsed, idx = [], []
    for i, c in enumerate(buf):
        if c == " " and collapsed and collapsed[-1] == " ":
            continue
        collapsed.append(c)
        idx.append(i)
    nc = "".join(collapsed)
    j = nc.find(nf)
    if j < 0:
        return None
    start = idx[j]
    end = idx[j + len(nf) - 1] + 1
    while end < len(text) and text[end] in " \t":
        end += 1
    return start


# ------------------------------------------------------------------ разметка --
TOKEN = re.compile(r"(\{-.*?-\}|\{\+.*?\+\})", re.S)

def parse_markup(text):
    segs = []
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("{-") and part.endswith("-}"):
            segs.append((part[2:-2], "red"))
        elif part.startswith("{+") and part.endswith("+}"):
            segs.append((part[2:-2], "green"))
        else:
            segs.append((part, ""))
    return segs

def build_new_from_src(text, repl, add):
    """Собирает новую редакцию статьи: замены помечаются зелёным, `add` дописывается в конец."""
    spans = []
    for find, put in repl:
        i = find_span(text, find)
        if i is None:
            raise ValueError(f"ЗАМЕНА: фрагмент не найден в статье:\n  ищем: {find[:150]!r}")
        spans.append((i, i + len(find), put))
    spans.sort(key=lambda x: x[0])
    for a, b in zip(spans, spans[1:]):
        if b[0] < a[1]:
            raise ValueError(f"ЗАМЕНА: пересечение фрагментов:\n  {a[2][:80]!r}\n  {b[2][:80]!r}")
    segs, pos = [], 0
    for s0, e0, put in spans:
        if s0 > pos:
            segs.append((text[pos:s0], ""))
        if put:
            segs.append((put, "green"))
        pos = e0
    if pos < len(text):
        segs.append((text[pos:], ""))
    if add:
        segs.append((("\n\n" if segs else "") + add, "green"))
    return segs


class R(dict):
    """Строка таблицы.

    src   = (акт, "номер статьи") — текст берётся дословно из файла пакета;
    red   = фрагменты, которые подсвечиваются красным в «настоящей редакции»;
    repl  = [(что ищем, на что меняем)] — замена помечается зелёным в «предложенном варианте»
            (если red не задан, красным помечаются те же фрагменты);
    add   = текст, дописываемый в конец статьи зелёным (новые части/примечания);
    old / new = полностью ручная разметка: {-красный-}, {+зелёный+}.
    """

    def __init__(self, src=None, red=None, repl=None, add=None, old=None, new=None,
                 note_old=None, note_new=None, red_all=False, green_all=False):
        self["src"] = src
        self["red_all"] = red_all
        self["green_all"] = green_all
        self["red"] = red if red is not None else ([f for f, _ in repl] if repl else [])
        self["repl"] = repl or []
        self["add"] = add
        self["old"] = old
        self["new"] = new
        self["note_old"] = note_old
        self["note_new"] = note_new

    def old_segs(self):
        if self["src"]:
            text = get_article(*self["src"])
            segs = [(text, "red")] if self["red_all"] else mark_red(text, self["red"])
        else:
            segs = parse_markup(self["old"])
        if self["note_old"]:
            segs = segs + [("\n\n" + self["note_old"], "note")]
        return segs

    def new_segs(self):
        if self["src"] and (self["repl"] or self["add"]):
            segs = build_new_from_src(get_article(*self["src"]), self["repl"], self["add"])
        else:
            segs = parse_markup(self["new"])
        if self["note_new"]:
            segs = segs + [("\n\n" + self["note_new"], "note")]
        return segs


# --------------------------------------------------------------------- docx ---
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x7A, 0x00)
GREY = RGBColor(0x40, 0x40, 0x40)
FONT = "Times New Roman"


def _set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _add_runs(par, segs, size=10):
    for text, kind in segs:
        parts = text.split("\n")
        for i, ptext in enumerate(parts):
            if i:
                par.add_run().add_break()
            if not ptext:
                continue
            run = par.add_run(ptext)
            run.font.name = FONT
            run.font.size = Pt(size)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            if kind == "red":
                run.font.color.rgb = RED
            elif kind == "green":
                run.font.color.rgb = GREEN
            elif kind == "note":
                run.font.color.rgb = GREY
                run.italic = True
                run.font.size = Pt(size - 1)


def _para(container, text, size=11, bold=False, align=None, space_after=6, italic=False):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def _page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    run.font.name = FONT
    run.font.size = Pt(9)


def build(doc, outdir=OUT):
    """doc: dict(act=..., title=..., date=..., essence=[..], section2=..., rows=[R..], notes=[..])"""
    d = Document()
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    _page_number_footer(sec)

    act = doc["act"]
    _para(d, f"ACT. {act}", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(d, "_" * 62, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(d, "IN THE SENATE OF THE STATE OF SAN ANDREAS", size=12, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(d, doc.get("date", "MARCH 00, 2026"), size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(d, "_" * 62, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _para(d, "ЗАКОНОПРОЕКТ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _para(d, "о внесении изменений в действующий закон", size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    _para(d, "SECTION 1. СУТЬ ПРАВКИ.", size=12, bold=True, space_after=4)
    for ab in doc["essence"]:
        _para(d, ab, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4)

    _para(d, f"SECTION 2. ВНЕСЕНИЕ ПРАВКИ В {doc['section2']}", size=12, bold=True, space_after=6)

    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    _repeat_header(hdr)
    usable_cm = sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm
    for i, name in enumerate(("настоящая редакция", "предложенный вариант")):
        cell = hdr.cells[i]
        cell.width = Cm(usable_cm / 2)
        _set_cell_bg(cell, "D9D9D9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(name)
        run.bold = True
        run.font.name = FONT
        run.font.size = Pt(11)

    for row in doc["rows"]:
        cells = table.add_row().cells
        for i, segs in enumerate((row.old_segs(), row.new_segs())):
            cells[i].width = Cm(usable_cm / 2)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            _add_runs(p, segs, size=10)

    if doc.get("notes"):
        _para(d, "", size=6, space_after=2)
        for i, n in enumerate(doc["notes"]):
            _para(d, n if i else "Примечания.", size=10, italic=(i == 0), bold=(i == 0),
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=3)

    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"ACT.{act}.docx")
    d.save(path)
    return path
